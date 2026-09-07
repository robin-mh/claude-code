# -*- coding: utf-8 -*-
"""Baut den re:school Zeiterfassungs-Tracker (Dashboard / Daten / Settings)."""

import datetime as dt

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.axis import ChartLines
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule, DataBarRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

OUT = "/home/user/claude-code/re-school-zeiterfassung.xlsx"

# ---------------------------------------------------------------- Farbsystem
DARK = "1D6F5E"   # re:school Petrol dunkel
MID = "4DA28E"    # re:school Petrol mittel
LIGHT = "CFE8E1"  # re:school Petrol hell
CARD_BG = "F5F9F8"
BAND = "EDF5F3"
EDGE = "C9D8D3"
TEXT = "333333"
LABEL = "7A8C88"
GREEN = "1F8A3B"
AMBER = "9A6A00"
RED = "C0392B"

FONT = "Arial"

f_dark = PatternFill("solid", fgColor=DARK)
f_mid = PatternFill("solid", fgColor=MID)
f_light = PatternFill("solid", fgColor=LIGHT)
f_card = PatternFill("solid", fgColor=CARD_BG)
f_band = PatternFill("solid", fgColor=BAND)
f_white = PatternFill("solid", fgColor="FFFFFF")

thin_edge = Side(style="thin", color=EDGE)
thin_dark = Side(style="thin", color=DARK)
med_mid = Side(style="medium", color=MID)
b_card = Border(left=thin_edge, right=thin_edge, top=thin_edge, bottom=thin_edge)
b_underline = Border(bottom=med_mid)
b_total = Border(top=thin_dark, bottom=Side(style="double", color=DARK))

# ---------------------------------------------------------------- Zahlenformate
NF_H = '#,##0.0" h"'
NF_H2 = '#,##0.00'
NF_CNT = '#,##0'
NF_PCT = '0.0%'
NF_SHARE = '0.0%" Anteil"'
NF_DELTA = '[Color10]0.0%"▲";[Red](0.0%)"▼";"–"'
NF_DELTA_ABS = '[Color10]#,##0" ▲";[Red](#,##0)" ▼";"–"'
NF_DATE = 'DD.MM.YYYY'
NF_TIME = 'HH:MM'
NF_HIDE0 = '#,##0.00;-#,##0.00;'
NF_INT0 = '0;-0;'

# ---------------------------------------------------------------- Stammdaten
PERSONEN = ["Robin", "Kathi"]
ARBEITSPAKETE = [
    "Konzept & Strategie",
    "Recherche",
    "Workshop-Vorbereitung",
    "Workshop-Durchführung",
    "Schulkontakt",
    "Material & Design",
    "Dokumentation",
    "Orga & Abstimmung",
    "Sonstiges",
]
STATUS = ["Geplant", "Läuft", "Erledigt"]
MONATE = ["Januar", "Februar", "März", "April", "Mai", "Juni",
          "Juli", "August", "September", "Oktober", "November", "Dezember"]
JAHRE = [2025, 2026, 2027]

LAST_ROW = 501  # Datenbereich Daten!2:501

# Beispielzeilen: Datum, Person, Arbeitspaket, Tätigkeit, Von, Bis, Pause(Min), Status
SAMPLE = [
    (dt.date(2026, 8, 4),  "Robin", "Konzept & Strategie",   "Zielbild re:school geschärft",        "09:00", "12:15", 15, "Erledigt"),
    (dt.date(2026, 8, 4),  "Kathi", "Recherche",             "Förderprogramme gesichtet",           "10:00", "12:30", 0,  "Erledigt"),
    (dt.date(2026, 8, 5),  "Kathi", "Schulkontakt",          "Erstgespräch GS Nordstadt",           "14:00", "15:30", 0,  "Erledigt"),
    (dt.date(2026, 8, 6),  "Robin", "Orga & Abstimmung",     "Wochenplanung mit Kathi",             "09:00", "10:00", 0,  "Erledigt"),
    (dt.date(2026, 8, 6),  "Kathi", "Orga & Abstimmung",     "Wochenplanung mit Robin",             "09:00", "10:00", 0,  "Erledigt"),
    (dt.date(2026, 8, 11), "Robin", "Workshop-Vorbereitung", "Ablauf Modul 1 entworfen",            "08:30", "13:00", 30, "Erledigt"),
    (dt.date(2026, 8, 12), "Kathi", "Material & Design",     "Kartenset Layout v1",                 "13:00", "17:45", 20, "Erledigt"),
    (dt.date(2026, 8, 13), "Robin", "Recherche",             "Benchmark Bildungsinitiativen",       "10:00", "12:00", 0,  "Erledigt"),
    (dt.date(2026, 8, 18), "Kathi", "Workshop-Vorbereitung", "Moderationsleitfaden Modul 1",        "09:15", "12:45", 15, "Erledigt"),
    (dt.date(2026, 8, 19), "Robin", "Schulkontakt",          "Terminabstimmung Pilotschulen",       "11:00", "12:30", 0,  "Erledigt"),
    (dt.date(2026, 8, 20), "Robin", "Material & Design",     "Slides Modul 1 gebaut",               "13:30", "18:00", 30, "Erledigt"),
    (dt.date(2026, 8, 25), "Kathi", "Workshop-Durchführung", "Pilot-Workshop GS Nordstadt",         "08:00", "13:00", 30, "Erledigt"),
    (dt.date(2026, 8, 25), "Robin", "Workshop-Durchführung", "Pilot-Workshop GS Nordstadt",         "08:00", "13:00", 30, "Erledigt"),
    (dt.date(2026, 8, 26), "Robin", "Dokumentation",         "Auswertung Pilot-Workshop",           "09:00", "11:30", 0,  "Erledigt"),
    (dt.date(2026, 8, 27), "Kathi", "Dokumentation",         "Feedbackbögen ausgewertet",           "14:00", "16:30", 0,  "Erledigt"),
    (dt.date(2026, 9, 1),  "Robin", "Konzept & Strategie",   "Modul 2 skizziert",                   "09:00", "12:00", 0,  "Erledigt"),
    (dt.date(2026, 9, 1),  "Kathi", "Orga & Abstimmung",     "Jour fixe + Backlog",                 "09:00", "10:15", 0,  "Erledigt"),
    (dt.date(2026, 9, 2),  "Kathi", "Material & Design",     "Kartenset v2 nach Feedback",          "10:00", "15:00", 45, "Erledigt"),
    (dt.date(2026, 9, 2),  "Robin", "Schulkontakt",          "Anfrage Sekundarstufe Ost",           "15:30", "16:30", 0,  "Erledigt"),
    (dt.date(2026, 9, 3),  "Robin", "Workshop-Vorbereitung", "Ablauf Modul 2",                      "08:45", "12:15", 15, "Erledigt"),
    (dt.date(2026, 9, 3),  "Kathi", "Recherche",             "Methodenpool erweitert",              "13:00", "15:30", 0,  "Erledigt"),
    (dt.date(2026, 9, 4),  "Robin", "Dokumentation",         "Projekttagebuch nachgezogen",         "16:00", "17:30", 0,  "Erledigt"),
    (dt.date(2026, 9, 7),  "Kathi", "Workshop-Vorbereitung", "Materialliste Modul 2",               "09:00", "11:00", 0,  "Läuft"),
    (dt.date(2026, 9, 7),  "Robin", "Orga & Abstimmung",     "Wochenplanung KW 37",                 "11:00", "12:00", 0,  "Läuft"),
]

wb = Workbook()

# ================================================================= DASHBOARD
dash = wb.active
dash.title = "Dashboard"
dash.sheet_view.showGridLines = False
dash.sheet_view.zoomScale = 90
dash.sheet_properties.tabColor = DARK

SLOTS = [("B", "C"), ("E", "F"), ("H", "I"), ("K", "L"), ("N", "O"), ("Q", "R")]
GUTTERS = ["D", "G", "J", "M", "P"]

widths = {"A": 2, "B": 13, "C": 11, "D": 2, "E": 11, "F": 11, "G": 2,
          "H": 11, "I": 11, "J": 2, "K": 13, "L": 11, "M": 2,
          "N": 11, "O": 11, "P": 2, "Q": 11, "R": 11, "S": 2}
for col, w in widths.items():
    dash.column_dimensions[col].width = w

for r, h in {1: 8, 2: 28, 3: 18, 4: 8, 5: 12, 6: 20, 7: 10,
             8: 4.5, 9: 13, 10: 30, 11: 15, 12: 10,
             13: 18, 14: 6, 15: 18, 26: 6, 27: 18, 28: 6}.items():
    dash.row_dimensions[r].height = h

# --- Banner
for r in (1, 2, 3):
    for c in range(1, 20):
        dash.cell(row=r, column=c).fill = f_dark
dash.merge_cells("B2:R2")
dash["B2"] = "re:school · Zeiterfassung"
dash["B2"].font = Font(name=FONT, size=20, bold=True, color="FFFFFF")
dash["B2"].alignment = Alignment(vertical="center")
dash.merge_cells("B3:R3")
dash["B3"] = "Stunden · Personen · Arbeitspakete"
dash["B3"].font = Font(name=FONT, size=10, color=LIGHT)
dash["B3"].alignment = Alignment(vertical="top")

# --- Filterzeile
FILTERS = [
    (0, "JAHR", "=Jahre", 2026, "0"),
    (1, "MONAT", "=F_Monat", "Alle", "General"),
    (2, "ARBEITSPAKET", "=F_AP", "Alle", "General"),
    (3, "STATUS", "=F_Status", "Alle", "General"),
]
for slot, label, src, default, nf in FILTERS:
    c1, c2 = SLOTS[slot]
    dash.merge_cells(f"{c1}5:{c2}5")
    dash[f"{c1}5"] = label
    dash[f"{c1}5"].font = Font(name=FONT, size=8, bold=True, color=LABEL)
    dash[f"{c1}5"].alignment = Alignment(horizontal="center", vertical="bottom")
    dash.merge_cells(f"{c1}6:{c2}6")
    cell = dash[f"{c1}6"]
    cell.value = default
    cell.number_format = nf
    cell.fill = f_light
    cell.font = Font(name=FONT, size=10, bold=True, color=DARK)
    cell.alignment = Alignment(horizontal="center", vertical="center")
    for col in (c1, c2):
        dash[f"{col}6"].border = Border(left=thin_dark, right=thin_dark,
                                        top=thin_dark, bottom=thin_dark)
    dv = DataValidation(type="list", formula1=src, allow_blank=False, showDropDown=False)
    dash.add_data_validation(dv)
    dv.add(dash[f"{c1}6"])

# --- KPI-Karten
J = "Daten!$J$2:$J$%d" % LAST_ROW   # Jahr
K = "Daten!$K$2:$K$%d" % LAST_ROW   # Monat
C = "Daten!$C$2:$C$%d" % LAST_ROW   # Arbeitspaket
I_ = "Daten!$I$2:$I$%d" % LAST_ROW  # Status
B = "Daten!$B$2:$B$%d" % LAST_ROW   # Person
H = "Daten!$H$2:$H$%d" % LAST_ROW   # Stunden

P_JAHR = "Settings!$C$16"
P_MONAT = "Settings!$C$17"
P_AP = "Settings!$C$18"
P_STATUS = "Settings!$C$19"
P_VMONAT = "Settings!$C$22"
P_VJAHR = "Settings!$C$23"


def sumifs(jahr, monat, extra=""):
    return (f"SUMIFS({H},{J},{jahr},{K},{monat},{C},{P_AP},{I_},{P_STATUS}{extra})")


def countifs(jahr, monat, extra=""):
    return (f"COUNTIFS({J},{jahr},{K},{monat},{C},{P_AP},{I_},{P_STATUS}{extra})")


CUR = sumifs(P_JAHR, P_MONAT)
PRE = sumifs(P_VJAHR, P_VMONAT)
CUR_N = countifs(P_JAHR, P_MONAT)
PRE_N = countifs(P_VJAHR, P_VMONAT)
ROBIN = sumifs(P_JAHR, P_MONAT, f',{B},"Robin"')
KATHI = sumifs(P_JAHR, P_MONAT, f',{B},"Kathi"')
OFFEN = (f'SUMIFS({H},{J},{P_JAHR},{K},{P_MONAT},{C},{P_AP},{I_},"<>Erledigt")')

CARDS = [
    ("GESAMTSTUNDEN", f"={CUR}", NF_H, f"=IFERROR({CUR}/{PRE}-1,0)", NF_DELTA),
    ("STUNDEN ROBIN", f"={ROBIN}", NF_H, f"=IFERROR({ROBIN}/{CUR},0)", NF_SHARE),
    ("STUNDEN KATHI", f"={KATHI}", NF_H, f"=IFERROR({KATHI}/{CUR},0)", NF_SHARE),
    ("EINTRÄGE", f"={CUR_N}", NF_CNT, f"={CUR_N}-{PRE_N}", NF_DELTA_ABS),
    ("Ø STUNDEN / EINTRAG", f"=IFERROR({CUR}/{CUR_N},0)", NF_H2,
     f"=IFERROR(IFERROR({CUR}/{CUR_N},0)/IFERROR({PRE}/{PRE_N},0)-1,0)", NF_DELTA),
    ("NOCH OFFEN", f"={OFFEN}", NF_H, f"=IFERROR({OFFEN}/{CUR},0)", NF_SHARE),
]

for idx, (label, val, nf, sub, sub_nf) in enumerate(CARDS):
    c1, c2 = SLOTS[idx]
    dash.merge_cells(f"{c1}8:{c2}8")
    dash.merge_cells(f"{c1}9:{c2}9")
    dash.merge_cells(f"{c1}10:{c2}10")
    dash.merge_cells(f"{c1}11:{c2}11")
    for col in (c1, c2):
        dash[f"{col}8"].fill = f_mid
    dash[f"{c1}9"] = label
    dash[f"{c1}9"].font = Font(name=FONT, size=8, bold=True, color=LABEL)
    dash[f"{c1}9"].alignment = Alignment(horizontal="center", vertical="center")
    dash[f"{c1}10"] = val
    dash[f"{c1}10"].number_format = nf
    dash[f"{c1}10"].font = Font(name=FONT, size=21, bold=True, color=DARK)
    dash[f"{c1}10"].alignment = Alignment(horizontal="center", vertical="center")
    dash[f"{c1}11"] = sub
    dash[f"{c1}11"].number_format = sub_nf
    dash[f"{c1}11"].font = Font(name=FONT, size=9, bold=True, color=TEXT)
    dash[f"{c1}11"].alignment = Alignment(horizontal="center", vertical="center")
    for row in (9, 10, 11):
        for col in (c1, c2):
            dash[f"{col}{row}"].fill = f_card
    # Perimeter
    for row in (8, 9, 10, 11):
        top = thin_edge if row == 8 else None
        bottom = thin_edge if row == 11 else None
        dash[f"{c1}{row}"].border = Border(left=thin_edge, top=top, bottom=bottom)
        dash[f"{c2}{row}"].border = Border(right=thin_edge, top=top, bottom=bottom)


def section(row, title):
    dash.merge_cells(f"B{row}:R{row}")
    dash[f"B{row}"] = title
    dash[f"B{row}"].font = Font(name=FONT, size=11, bold=True, color=DARK)
    dash[f"B{row}"].alignment = Alignment(vertical="center")
    for col in range(2, 19):
        dash.cell(row=row, column=col).border = b_underline


section(13, "AUFTEILUNG")
section(27, "VERLAUF & VERTEILUNG")


def table_header(row, cols, titles, first_left=True):
    """cols: Liste von (start, end) Spaltenpaaren; füllt auch die Gutter dazwischen."""
    start_col = cols[0][0]
    end_col = cols[-1][1]
    for c in range(ord(start_col) - 64, ord(end_col) - 64 + 1):
        cell = dash.cell(row=row, column=c)
        cell.fill = f_dark
    for i, ((c1, c2), title) in enumerate(zip(cols, titles)):
        dash.merge_cells(f"{c1}{row}:{c2}{row}")
        dash[f"{c1}{row}"] = title
        dash[f"{c1}{row}"].font = Font(name=FONT, size=10, bold=True, color="FFFFFF")
        align = "left" if (i == 0 and first_left) else "right"
        dash[f"{c1}{row}"].alignment = Alignment(horizontal=align, vertical="center")


def table_row(row, cols, values, formats, aligns, band, bold=False, border=None):
    start_col = cols[0][0]
    end_col = cols[-1][1]
    for c in range(ord(start_col) - 64, ord(end_col) - 64 + 1):
        cell = dash.cell(row=row, column=c)
        if band:
            cell.fill = f_band
        if border:
            cell.border = border
    for (c1, c2), val, nf, al in zip(cols, values, formats, aligns):
        dash.merge_cells(f"{c1}{row}:{c2}{row}")
        cell = dash[f"{c1}{row}"]
        cell.value = val
        cell.number_format = nf
        cell.font = Font(name=FONT, size=10, bold=bold,
                         color=DARK if bold else TEXT)
        cell.alignment = Alignment(horizontal=al, vertical="center")


# --- Tabelle 1: Stunden je Person (Slots 1-3)
T1 = [("B", "C"), ("E", "F"), ("H", "I")]
table_header(15, T1, ["PERSON", "STUNDEN", "Δ VORPERIODE"])
for i, person in enumerate(PERSONEN):
    r = 16 + i
    cur = sumifs(P_JAHR, P_MONAT, f',{B},"{person}"')
    pre = sumifs(P_VJAHR, P_VMONAT, f',{B},"{person}"')
    table_row(r, T1, [person, f"={cur}", f"=IFERROR({cur}/{pre}-1,0)"],
              ["General", NF_H2, NF_DELTA], ["left", "right", "right"], band=(i % 2 == 0))
table_row(18, T1, ["Gesamt", "=SUM(E16:E17)", f"=IFERROR({CUR}/{PRE}-1,0)"],
          ["General", NF_H2, NF_DELTA], ["left", "right", "right"],
          band=False, bold=True, border=b_total)

# --- Tabelle 2: Stunden je Arbeitspaket (Slots 4-6)
T2 = [("K", "L"), ("N", "O"), ("Q", "R")]
table_header(15, T2, ["ARBEITSPAKET", "STUNDEN", "ANTEIL"])
for i, ap in enumerate(ARBEITSPAKETE):
    r = 16 + i
    cur = sumifs(P_JAHR, P_MONAT, f',{C},"{ap}"')
    table_row(r, T2, [ap, f"={cur}", f"=IFERROR({cur}/$N$25,0)"],
              ["General", NF_H2, NF_PCT], ["left", "right", "right"], band=(i % 2 == 0))
table_row(25, T2, ["Gesamt", "=SUM(N16:N24)", "=IFERROR(N25/N25,0)"],
          ["General", NF_H2, NF_PCT], ["left", "right", "right"],
          band=False, bold=True, border=b_total)

# Datenbalken auf den Stunden-Spalten
dash.conditional_formatting.add(
    "E16:E17", DataBarRule(start_type="num", start_value=0,
                           end_type="max", color=MID, showValue=True))
dash.conditional_formatting.add(
    "N16:N24", DataBarRule(start_type="num", start_value=0,
                           end_type="max", color=MID, showValue=True))

dash.freeze_panes = "A13"
dash.print_area = "A1:S40"
dash.page_setup.orientation = "landscape"
dash.page_setup.fitToWidth = 1
dash.page_setup.fitToHeight = 1
dash.sheet_properties.pageSetUpPr.fitToPage = True

# ================================================================= DATEN
data = wb.create_sheet("Daten")
data.sheet_view.showGridLines = False
data.sheet_properties.tabColor = MID

HEADERS = ["Datum", "Person", "Arbeitspaket", "Tätigkeit", "Von", "Bis",
           "Pause (Min)", "Stunden", "Status", "Jahr", "Monat", "KW"]
COLW = [13, 12, 24, 38, 8, 8, 11, 11, 12, 9, 13, 7]
for i, (h, w) in enumerate(zip(HEADERS, COLW), start=1):
    cell = data.cell(row=1, column=i, value=h)
    cell.fill = f_dark
    cell.font = Font(name=FONT, size=10, bold=True, color="FFFFFF")
    if i in (1, 9, 10, 12):
        h_align = "center"
    elif i in (2, 3, 4, 11):
        h_align = "left"
    else:
        h_align = "right"
    cell.alignment = Alignment(horizontal=h_align, vertical="center")
    data.column_dimensions[get_column_letter(i)].width = w
data.row_dimensions[1].height = 20

COMMENTS = {
    "A1": "Eingabefeld: Datum des Eintrags (TT.MM.JJJJ).",
    "B1": "Eingabefeld: Auswahl Robin / Kathi (Liste auf 'Settings').",
    "C1": "Eingabefeld: Arbeitspaket aus der Auswahlliste. Pakete auf 'Settings' anpassbar.",
    "D1": "Eingabefeld: kurze Beschreibung der Tätigkeit (Freitext).",
    "E1": "Eingabefeld: Startzeit, Format HH:MM.",
    "F1": "Eingabefeld: Endzeit, Format HH:MM. Über Mitternacht wird korrekt gerechnet.",
    "G1": "Eingabefeld: Pausenzeit in Minuten (leer = 0).",
    "H1": "Berechnet: (Bis - Von) - Pause, in Dezimalstunden. Nicht überschreiben.",
    "I1": "Eingabefeld: Geplant / Läuft / Erledigt.",
    "J1": "Hilfsspalte, berechnet aus dem Datum. Nicht überschreiben.",
    "K1": "Hilfsspalte, berechnet aus dem Datum. Nicht überschreiben.",
    "L1": "Hilfsspalte: Kalenderwoche (Woche beginnt Montag). Nicht überschreiben.",
}
for ref, txt in COMMENTS.items():
    data[ref].comment = Comment(txt, "re:school")

data["A1"].comment = Comment(
    "Eingabefeld: Datum des Eintrags (TT.MM.JJJJ).\n\n"
    "Ausfüllen: Spalten A-G und I. Die Spalten H, J, K, L rechnen sich selbst.\n"
    "Die Zeilen 2-25 sind Beispieldaten - einfach überschreiben oder löschen.",
    "re:school")

for r in range(2, LAST_ROW + 1):
    idx = r - 2
    if idx < len(SAMPLE):
        d, person, ap, tat, von, bis, pause, status = SAMPLE[idx]
        data.cell(row=r, column=1, value=d)
        data.cell(row=r, column=2, value=person)
        data.cell(row=r, column=3, value=ap)
        data.cell(row=r, column=4, value=tat)
        data.cell(row=r, column=5, value=dt.time(*map(int, von.split(":"))))
        data.cell(row=r, column=6, value=dt.time(*map(int, bis.split(":"))))
        data.cell(row=r, column=7, value=pause)
        data.cell(row=r, column=9, value=status)
    data.cell(row=r, column=8,
              value=f'=IF(OR($E{r}="",$F{r}=""),0,'
                    f'ROUND(MOD($F{r}-$E{r},1)*24-N($G{r})/60,2))')
    data.cell(row=r, column=10, value=f'=IF($A{r}="",0,YEAR($A{r}))')
    data.cell(row=r, column=11, value=f'=IF($A{r}="","",INDEX(Monate,MONTH($A{r})))')
    data.cell(row=r, column=12, value=f'=IF($A{r}="",0,WEEKNUM($A{r},2))')

    band = (r % 2 == 0)
    for c in range(1, 13):
        cell = data.cell(row=r, column=c)
        cell.font = Font(name=FONT, size=10, color=TEXT)
        if band:
            cell.fill = f_band
    data.cell(row=r, column=1).number_format = NF_DATE
    data.cell(row=r, column=5).number_format = NF_TIME
    data.cell(row=r, column=6).number_format = NF_TIME
    data.cell(row=r, column=7).number_format = NF_INT0
    data.cell(row=r, column=8).number_format = NF_HIDE0
    data.cell(row=r, column=8).font = Font(name=FONT, size=10, bold=True, color=DARK)
    data.cell(row=r, column=10).number_format = NF_INT0
    data.cell(row=r, column=12).number_format = NF_INT0
    for c in (7, 8):
        data.cell(row=r, column=c).alignment = Alignment(horizontal="right")
    for c in (1, 5, 6, 9, 10, 12):
        data.cell(row=r, column=c).alignment = Alignment(horizontal="center")

rng = f"2:{LAST_ROW}"
dv_person = DataValidation(type="list", formula1="=Personen", allow_blank=True, showDropDown=False)
dv_ap = DataValidation(type="list", formula1="=Arbeitspakete", allow_blank=True, showDropDown=False)
dv_status = DataValidation(type="list", formula1="=StatusListe", allow_blank=True, showDropDown=False)
for dv, col in ((dv_person, "B"), (dv_ap, "C"), (dv_status, "I")):
    data.add_data_validation(dv)
    dv.add(f"{col}2:{col}{LAST_ROW}")

status_colors = {"Erledigt": GREEN, "Läuft": AMBER, "Geplant": LABEL}
for value, color in status_colors.items():
    data.conditional_formatting.add(
        f"I2:I{LAST_ROW}",
        CellIsRule(operator="equal", formula=[f'"{value}"'],
                   font=Font(name=FONT, size=10, bold=True, color=color)))

data.freeze_panes = "B2"
data.print_title_rows = "1:1"
data.page_setup.orientation = "landscape"
data.page_setup.fitToWidth = 1
data.page_setup.fitToHeight = 0
data.sheet_properties.pageSetUpPr.fitToPage = True
data.auto_filter.ref = f"A1:L{LAST_ROW}"

# ================================================================= SETTINGS
st = wb.create_sheet("Settings")
st.sheet_view.showGridLines = False
st.sheet_properties.tabColor = LIGHT

for col, w in {"A": 2, "B": 22, "C": 22, "D": 15, "E": 2, "F": 24, "G": 12,
               "H": 14, "I": 2, "J": 10, "K": 2, "L": 14, "M": 2,
               "N": 24, "O": 2, "P": 12}.items():
    st.column_dimensions[col].width = w


def st_header(ref, title, span=1):
    col = ref[0]
    row = int(ref[1:])
    end = get_column_letter(ord(col) - 64 + span - 1)
    if span > 1:
        st.merge_cells(f"{col}{row}:{end}{row}")
    st[ref] = title
    for c in range(ord(col) - 64, ord(col) - 64 + span):
        cell = st.cell(row=row, column=c)
        cell.fill = f_dark
        cell.font = Font(name=FONT, size=9, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="left", vertical="center")


def st_list(col, start, values, nf="General"):
    for i, v in enumerate(values):
        cell = st[f"{col}{start + i}"]
        cell.value = v
        cell.number_format = nf
        cell.font = Font(name=FONT, size=10, color=TEXT)
        cell.border = Border(bottom=Side(style="hair", color=EDGE))


st_header("B2", "PERSONEN")
st_list("B", 3, PERSONEN)
st_header("F2", "ARBEITSPAKETE")
st_list("F", 3, ARBEITSPAKETE)
st_header("H2", "STATUS")
st_list("H", 3, STATUS)
st_header("J2", "JAHRE")
st_list("J", 3, JAHRE, "0")
st_header("L2", "MONATE")
st_list("L", 3, MONATE)
st_header("N2", "FILTER: MONAT")
st_list("N", 3, ["Alle"] + MONATE)
st_list("P", 3, ["Alle"] + ARBEITSPAKETE)
st.column_dimensions["P"].width = 24
st_header("P2", "FILTER: ARBEITSPAKET")
st_list("D", 3, ["Alle"] + STATUS)
st_header("D2", "FILTER: STATUS")

# --- Parameter / Vorperioden-Helfer
st_header("B15", "PARAMETER & VORPERIODE", span=2)
params = [
    ("Jahr (Auswahl)", "=Dashboard!$B$6", "0"),
    ("Monat-Kriterium", '=IF(Dashboard!$E$6="Alle","*",Dashboard!$E$6)', "General"),
    ("Arbeitspaket-Kriterium", '=IF(Dashboard!$H$6="Alle","*",Dashboard!$H$6)', "General"),
    ("Status-Kriterium", '=IF(Dashboard!$K$6="Alle","*",Dashboard!$K$6)', "General"),
    ("Monat-Index", "=IFERROR(MATCH(Dashboard!$E$6,Monate,0),0)", NF_CNT),
    ("Vormonat-Index", "=IF($C$20=0,0,IF($C$20=1,12,$C$20-1))", NF_CNT),
    ("Vormonat-Kriterium", '=IF($C$21=0,"*",INDEX(Monate,$C$21))', "General"),
    ("Vorperioden-Jahr", "=IF($C$20=0,$C$16-1,IF($C$20=1,$C$16-1,$C$16))", "0"),
]
for i, (label, formula, nf) in enumerate(params):
    r = 16 + i
    st[f"B{r}"] = label
    st[f"B{r}"].font = Font(name=FONT, size=9, color=LABEL)
    st[f"C{r}"] = formula
    st[f"C{r}"].number_format = nf
    st[f"C{r}"].font = Font(name=FONT, size=10, bold=True, color=DARK)
    st[f"C{r}"].alignment = Alignment(horizontal="right")
    for c in ("B", "C"):
        st[f"{c}{r}"].border = Border(bottom=Side(style="hair", color=EDGE))
st["C16"].comment = Comment(
    "Vorperiode = Vormonat (bei Monatsauswahl) bzw. Vorjahr (bei 'Alle').", "re:school")

# --- Chart-Hilfstabellen
st_header("B27", "MONAT")
st_header("C27", "STUNDEN")
st_header("D27", "EINTRÄGE")
for i, m in enumerate(MONATE):
    r = 28 + i
    st[f"B{r}"] = m
    st[f"B{r}"].font = Font(name=FONT, size=9, color=TEXT)
    st[f"C{r}"] = (f'=SUMIFS({H},{J},$C$16,{K},$B{r},{C},$C$18,{I_},$C$19)')
    st[f"D{r}"] = (f'=COUNTIFS({J},$C$16,{K},$B{r},{C},$C$18,{I_},$C$19)')
    st[f"C{r}"].number_format = NF_H2
    st[f"D{r}"].number_format = NF_CNT
    for c in ("C", "D"):
        st[f"{c}{r}"].font = Font(name=FONT, size=9, color=TEXT)

st_header("F27", "ARBEITSPAKET")
st_header("G27", "STUNDEN")
for i, ap in enumerate(ARBEITSPAKETE):
    r = 28 + i
    st[f"F{r}"] = ap
    st[f"F{r}"].font = Font(name=FONT, size=9, color=TEXT)
    st[f"G{r}"] = (f'=SUMIFS({H},{J},$C$16,{K},$C$17,{C},$F{r},{I_},$C$19)')
    st[f"G{r}"].number_format = NF_H2
    st[f"G{r}"].font = Font(name=FONT, size=9, color=TEXT)

# --- Akzent-Referenz
st_header("B42", "AKZENTFARBEN", span=2)
for i, (name, hexv, fill) in enumerate([("Dunkel", DARK, f_dark),
                                        ("Mittel", MID, f_mid),
                                        ("Hell", LIGHT, f_light)]):
    r = 43 + i
    st[f"B{r}"] = name
    st[f"B{r}"].font = Font(name=FONT, size=9, color=LABEL)
    st[f"C{r}"] = f"#{hexv}"
    st[f"C{r}"].fill = fill
    st[f"C{r}"].font = Font(name=FONT, size=9, bold=True,
                            color="FFFFFF" if i < 2 else DARK)
    st[f"C{r}"].alignment = Alignment(horizontal="center")

# ================================================================= Namen
for name, ref in [
    ("Personen", "Settings!$B$3:$B$4"),
    ("Arbeitspakete", f"Settings!$F$3:$F${2 + len(ARBEITSPAKETE)}"),
    ("StatusListe", "Settings!$H$3:$H$5"),
    ("Jahre", "Settings!$J$3:$J$5"),
    ("Monate", "Settings!$L$3:$L$14"),
    ("F_Monat", "Settings!$N$3:$N$15"),
    ("F_AP", f"Settings!$P$3:$P${3 + len(ARBEITSPAKETE)}"),
    ("F_Status", "Settings!$D$3:$D$6"),
]:
    wb.defined_names.add(DefinedName(name, attr_text=ref))

# ================================================================= Charts
trend = BarChart()
trend.type = "col"
trend.style = None
trend.title = "Stunden je Monat"
trend.height = 8.4
trend.width = 13.4
d_hours = Reference(st, min_col=3, min_row=27, max_row=39)
cats = Reference(st, min_col=2, min_row=28, max_row=39)
trend.add_data(d_hours, titles_from_data=True)
trend.set_categories(cats)
trend.y_axis.title = "Stunden"
trend.gapWidth = 55
trend.series[0].graphicalProperties.solidFill = DARK
trend.series[0].graphicalProperties.line.noFill = True

line = LineChart()
d_cnt = Reference(st, min_col=4, min_row=27, max_row=39)
line.add_data(d_cnt, titles_from_data=True)
line.y_axis.axId = 200
line.y_axis.title = "Einträge"
line.series[0].graphicalProperties.line.solidFill = MID
line.series[0].graphicalProperties.line.width = 22000
line.series[0].smooth = False
trend.y_axis.crosses = "autoZero"
line.y_axis.crosses = "max"
trend += line

for ch in (trend,):
    ch.x_axis.delete = False
    ch.y_axis.delete = False
    ch.y_axis.majorGridlines = ChartLines()

dist = BarChart()
dist.type = "bar"
dist.style = None
dist.title = "Stunden je Arbeitspaket"
dist.height = 8.4
dist.width = 13.4
d_ap = Reference(st, min_col=7, min_row=27, max_row=27 + len(ARBEITSPAKETE))
cats_ap = Reference(st, min_col=6, min_row=28, max_row=27 + len(ARBEITSPAKETE))
dist.add_data(d_ap, titles_from_data=True)
dist.set_categories(cats_ap)
dist.gapWidth = 45
dist.legend = None
dist.series[0].graphicalProperties.solidFill = MID
dist.series[0].graphicalProperties.line.noFill = True
dist.x_axis.delete = False
dist.y_axis.delete = False

dash.add_chart(trend, "B29")
dash.add_chart(dist, "K29")

wb.active = 0
wb.save(OUT)
print("gespeichert:", OUT)
